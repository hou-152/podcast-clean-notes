#!/usr/bin/env python3
"""Build the Feishu-import DOCX for course-feishu-pack.

Combines 02 三级笔记, 03 概念提取, and the untouched source (txt transcript or
original docx) into `{主题}_三级笔记_概念提取_原文_飞书导入版.docx`, with
self-verification (hashes, body suffix, media, relationships) and a manifest.json.

Usage:
  python3 build_feishu_docx.py \
    --title "8月29日 第二课" \
    --notes "02_主题_三级笔记.md" \
    --concepts "03_主题_概念提取.md" \
    --source "00_原始转写.txt" \
    --output "主题_三级笔记_概念提取_原文_飞书导入版.docx" \
    [--subtitle "课程整理·2026-08-31"] \
    [--assets-dir "assets"]
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Arial"
FONT_CJK = "Noto Sans CJK SC"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(85, 85, 85)
LIGHT_FILL = "F5F5F5"
STYLE_PREFIX = "FeishuPack"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")
IMAGE_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)$")
HEADING_RE = re.compile(r"^(#{1,4})\s+(.+)$")
BULLET_RE = re.compile(r"^[-*]\s+(.+)$")
ORDERED_RE = re.compile(r"^\d+[.)]\s+(.+)$")
SPEAKER_RE = re.compile(r"^(.+?)\s{2,}(\d{1,2}:\d{2}(?::\d{2})?)$")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_run_font(run, *, size: float, bold: bool = False, italic: bool = False,
                 color: RGBColor = BLACK) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT_CJK if attr == "eastAsia" else FONT)


def new_style(doc: Document, key: str, *, size: float, color: RGBColor = BLACK,
              bold: bool = False, italic: bool = False, before: float = 0,
              after: float = 8, line_spacing: float = 1.15,
              keep: bool = False, outline: int | None = None) -> str:
    name = f"{STYLE_PREFIX}{key}"
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT_CJK if attr == "eastAsia" else FONT)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    pf.keep_with_next = keep
    pf.widow_control = True
    if outline is not None:
        ppr = style.element.get_or_add_pPr()
        level = OxmlElement("w:outlineLvl")
        level.set(qn("w:val"), str(outline))
        ppr.append(level)
    return name


def set_fill(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _numbering_element(doc: Document):
    try:
        return doc.part.numbering_part.element
    except NotImplementedError:
        # Source DOCX has no word/numbering.xml (e.g. list-free docx sources);
        # python-docx cannot create one, so build a minimal empty part here.
        from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import XmlPart
        part = XmlPart(
            PackURI("/word/numbering.xml"),
            CT.WML_NUMBERING,
            parse_xml(
                '<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
            ),
            doc.part.package,
        )
        doc.part.relate_to(part, RT.NUMBERING)
        return part.element


def add_numbering(doc: Document, *, ordered: bool, left: int, hanging: int) -> int:
    numbering = _numbering_element(doc)
    abstract_id = max(
        [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))],
        default=0,
    ) + 1
    num_id = max(
        [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))],
        default=0,
    ) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"),
                     ("w:numFmt", "decimal" if ordered else "bullet"),
                     ("w:lvlText", "%1." if ordered else "●"),
                     ("w:suff", "tab")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        lvl.append(el)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))
    ppr.append(num_pr)
    if level:
        paragraph.paragraph_format.left_indent = Inches(0.25 * level)


def add_inline(paragraph, text: str, *, size: float = 11,
               color: RGBColor = BLACK, italic: bool = False) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, italic=italic, color=color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size - 0.5, italic=italic, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, italic=italic, color=color)


def iter_markdown_blocks(path: Path, assets_dir: Path | None = None):
    """Yield (kind, level, payload) for a markdown file.

    Image targets resolve against the markdown's own directory first, then
    assets_dir. Mermaid fenced blocks are skipped (Feishu cannot render them);
    their surrounding prose stays.
    """
    in_fence = False
    fence_language = ""
    base = path.parent
    for raw in path.read_text(encoding="utf-8").splitlines():
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_fence:
                in_fence = False
                fence_language = ""
            else:
                in_fence = True
                fence_language = stripped[3:].strip().lower()
            continue
        if not stripped:
            continue
        if in_fence:
            if fence_language == "mermaid":
                continue
            yield "code", 0, stripped
            continue
        heading = HEADING_RE.match(stripped)
        if heading:
            yield "heading", len(heading.group(1)), heading.group(2)
            continue
        image = IMAGE_RE.match(stripped)
        if image:
            alt, target = image.group(1), image.group(2)
            cand = (base / target).resolve()
            if not cand.is_file() and assets_dir is not None:
                cand = (assets_dir / target).resolve()
            if not cand.is_file():
                tried = f"tried {base}"
                if assets_dir is not None:
                    tried += f" and {assets_dir}"
                raise FileNotFoundError(f"Markdown image not found: {target} ({tried})")
            yield "image", 0, (alt, str(cand))
            continue
        if stripped.startswith(">"):
            quote_text = re.sub(r"^(?:>\s*)+", "", stripped)
            if quote_text:
                yield "quote", 0, quote_text
            continue
        bullet = BULLET_RE.match(stripped)
        if bullet:
            yield "bullet", max(0, (len(raw) - len(raw.lstrip())) // 2), bullet.group(1)
            continue
        ordered = ORDERED_RE.match(stripped)
        if ordered:
            yield "ordered", max(0, (len(raw) - len(raw.lstrip())) // 2), ordered.group(1)
            continue
        yield "body", 0, stripped


def iter_source_lines(source: Path):
    """Yield (kind, text) for a transcript txt: 'speaker' or 'body'."""
    for raw in source.read_text(encoding="utf-8").splitlines():
        stripped = raw.strip()
        if not stripped:
            continue
        if SPEAKER_RE.match(stripped):
            yield "speaker", stripped
        else:
            yield "body", stripped


def add_markdown(doc: Document, path: Path, styles: dict,
                 bullet_num: int, ordered_num: int, assets_dir: Path | None = None):
    result = []
    heading_sizes = {1: 20, 2: 16, 3: 14, 4: 12.5}
    for kind, level, text in iter_markdown_blocks(path, assets_dir):
        if kind == "heading":
            paragraph = doc.add_paragraph(style=styles[f"h{min(level, 4)}"])
            add_inline(paragraph, text, size=heading_sizes[min(level, 4)])
        elif kind == "quote":
            paragraph = doc.add_paragraph(style=styles["quote"])
            add_inline(paragraph, text, size=10.5, color=GRAY, italic=True)
        elif kind == "code":
            paragraph = doc.add_paragraph(style=styles["code"])
            add_inline(paragraph, text, size=9.5, color=GRAY)
        elif kind == "image":
            alt, image_path = text
            paragraph = doc.add_paragraph(style=styles["figure"])
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(image_path, width=Inches(5.95))
            doc_pr = run._element.xpath('.//*[local-name()="docPr"]')
            if doc_pr:
                doc_pr[0].set("descr", alt)
                doc_pr[0].set("title", alt.split("｜", 1)[0])
            caption = doc.add_paragraph(style=styles["caption"])
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(caption, alt, size=9.5, color=GRAY)
        elif kind in ("bullet", "ordered"):
            paragraph = doc.add_paragraph(style=styles["body"])
            apply_num(paragraph, bullet_num if kind == "bullet" else ordered_num, level)
            add_inline(paragraph, text)
        else:
            paragraph = doc.add_paragraph(style=styles["body"])
            add_inline(paragraph, text)
        result.append(paragraph)
    return result


def add_source_txt(doc: Document, source: Path, styles: dict, title: str) -> None:
    p = doc.add_paragraph(style=styles["part"])
    p.paragraph_format.page_break_before = True
    add_inline(p, "第三部分｜原始转写（未改写）", size=20)
    note = doc.add_paragraph(style=styles["quote"])
    add_inline(
        note,
        f"以下内容完整保留自 {source.name}；其中的 ASR 错字、时间戳与口语原声均未改写。",
        size=10.5, color=GRAY, italic=True)
    for kind, text in iter_source_lines(source):
        if kind == "speaker":
            p = doc.add_paragraph(style=styles["speaker"])
            add_inline(p, text, size=10, color=GRAY, italic=False)
        else:
            p = doc.add_paragraph(style=styles["body"])
            add_inline(p, text, size=11)


def c14n_blocks(path: Path):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    body = root.find("w:body", ns)
    return [
        etree.tostring(child, method="c14n")
        for child in body
        if etree.QName(child).localname != "sectPr"
    ]


def package_manifest(path: Path, prefix: str) -> dict[str, str]:
    with ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
            if name.startswith(prefix) and not name.endswith("/")
        }


def relationship_manifest(path: Path) -> dict[str, tuple[str, str]]:
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
    return {
        el.get("Id"): (el.get("Type"), el.get("Target"))
        for el in root
    }


def ensure_page_setup(doc: Document) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    if sect_pr is None:
        sect_pr = OxmlElement("w:sectPr")
        body.append(sect_pr)
    page_size = sect_pr.find(qn("w:pgSz"))
    if page_size is None:
        page_size = OxmlElement("w:pgSz")
        sect_pr.append(page_size)
    page_size.set(qn("w:w"), "12240")
    page_size.set(qn("w:h"), "15840")
    page_margins = sect_pr.find(qn("w:pgMar"))
    if page_margins is None:
        page_margins = OxmlElement("w:pgMar")
        sect_pr.append(page_margins)
    for side in ("top", "right", "bottom", "left"):
        page_margins.set(qn(f"w:{side}"), "1440")
    page_margins.set(qn("w:header"), "708")
    page_margins.set(qn("w:footer"), "708")
    page_margins.set(qn("w:gutter"), "0")


def build_styles(doc: Document) -> dict:
    styles = {
        "title": new_style(doc, "Title", size=26, after=3, line_spacing=1.15, keep=True),
        "subtitle": new_style(doc, "Subtitle", size=10.5, italic=True, color=GRAY, after=8),
        "part": new_style(doc, "Part", size=20, before=20, after=6, keep=True, outline=0),
        "h1": new_style(doc, "H1", size=20, before=20, after=6, keep=True, outline=1),
        "h2": new_style(doc, "H2", size=16, before=18, after=6, keep=True, outline=2),
        "h3": new_style(doc, "H3", size=14, color=GRAY, before=16, after=4, keep=True, outline=3),
        "h4": new_style(doc, "H4", size=12.5, color=GRAY, before=12, after=3, keep=True, outline=4),
        "body": new_style(doc, "Body", size=11, after=8, line_spacing=1.15),
        "quote": new_style(doc, "Quote", size=10.5, italic=True, color=GRAY, after=8, line_spacing=1.15),
        "code": new_style(doc, "Code", size=9.5, color=GRAY, after=4, line_spacing=1.0),
        "figure": new_style(doc, "Figure", size=9.5, after=3, line_spacing=1.0, keep=True),
        "caption": new_style(doc, "Caption", size=9.5, italic=True, color=GRAY, after=10, line_spacing=1.0),
        "speaker": new_style(doc, "Speaker", size=10, color=GRAY, before=6, after=2, line_spacing=1.0, keep=True),
    }
    doc.styles[styles["quote"]].paragraph_format.left_indent = Inches(0.3)
    doc.styles[styles["quote"]].paragraph_format.right_indent = Inches(0.15)
    return styles


def build_txt(source: Path, notes: Path, concepts: Path, output: Path,
              title: str, subtitle: str, assets_dir: Path | None) -> dict:
    hashes = {
        "source": sha256(source),
        "notes": sha256(notes),
        "concepts": sha256(concepts),
    }
    doc = Document()
    styles = build_styles(doc)
    bullet_num = add_numbering(doc, ordered=False, left=720, hanging=360)
    ordered_num = add_numbering(doc, ordered=True, left=720, hanging=360)

    p = doc.add_paragraph(style=styles["title"])
    add_inline(p, title, size=26)
    if subtitle:
        p = doc.add_paragraph(style=styles["subtitle"])
        add_inline(p, subtitle, size=10.5, color=GRAY, italic=True)
    info = doc.add_paragraph(style=styles["body"])
    add_inline(info, "文档顺序：三级笔记 → 概念提取 → 原始转写（未改写）。")
    set_fill(info, LIGHT_FILL)
    info.paragraph_format.left_indent = Inches(0.15)
    info.paragraph_format.right_indent = Inches(0.15)
    info.paragraph_format.space_before = Pt(6)
    info.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph(style=styles["part"])
    add_inline(p, "第一部分｜三级笔记", size=20)
    add_markdown(doc, notes, styles, bullet_num, ordered_num, assets_dir)

    p = doc.add_paragraph(style=styles["part"])
    p.paragraph_format.page_break_before = True
    add_inline(p, "第二部分｜概念提取", size=20)
    add_markdown(doc, concepts, styles, bullet_num, ordered_num, assets_dir)

    add_source_txt(doc, source, styles, title)

    ensure_page_setup(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    hashes["output"] = sha256(output)

    print(f"source_sha256={hashes['source']}")
    print(f"notes_sha256={hashes['notes']}")
    print(f"concepts_sha256={hashes['concepts']}")
    print(f"output_sha256={hashes['output']}")
    return hashes


def build_docx(source: Path, notes: Path, concepts: Path, output: Path,
               title: str, subtitle: str, assets_dir: Path | None) -> dict:
    source_hash_before = sha256(source)
    source_blocks = c14n_blocks(source)
    source_media = package_manifest(source, "word/media/")
    source_rels = relationship_manifest(source)

    doc = Document(source)
    body = doc._element.body
    if len(body) == 0:
        raise RuntimeError("Source DOCX has no body anchor")
    anchor = body[0]
    styles = build_styles(doc)
    bullet_num = add_numbering(doc, ordered=False, left=720, hanging=360)
    ordered_num = add_numbering(doc, ordered=True, left=720, hanging=360)

    inserted = []
    p = doc.add_paragraph(style=styles["title"])
    add_inline(p, title, size=26)
    inserted.append(p)
    if subtitle:
        p = doc.add_paragraph(style=styles["subtitle"])
        add_inline(p, subtitle, size=10.5, color=GRAY, italic=True)
        inserted.append(p)
    info = doc.add_paragraph(style=styles["body"])
    add_inline(info, "文档顺序：三级笔记 → 概念提取 → 用户提供的原始 DOCX。原文部分未作改写，原有图片完整保留。")
    set_fill(info, LIGHT_FILL)
    info.paragraph_format.left_indent = Inches(0.15)
    info.paragraph_format.right_indent = Inches(0.15)
    info.paragraph_format.space_before = Pt(6)
    info.paragraph_format.space_after = Pt(14)
    inserted.append(info)

    p = doc.add_paragraph(style=styles["part"])
    add_inline(p, "第一部分｜三级笔记", size=20)
    inserted.append(p)
    inserted.extend(add_markdown(doc, notes, styles, bullet_num, ordered_num, assets_dir))

    p = doc.add_paragraph(style=styles["part"])
    p.paragraph_format.page_break_before = True
    add_inline(p, "第二部分｜概念提取", size=20)
    inserted.append(p)
    inserted.extend(add_markdown(doc, concepts, styles, bullet_num, ordered_num, assets_dir))

    p = doc.add_paragraph(style=styles["part"])
    p.paragraph_format.page_break_before = True
    add_inline(p, "第三部分｜原始文档（未改写）", size=20)
    inserted.append(p)
    note = doc.add_paragraph(style=styles["quote"])
    add_inline(note, "以下内容完整保留自用户提供的原始 DOCX，未作改写。", size=10.5, color=GRAY, italic=True)
    inserted.append(note)

    for paragraph in inserted:
        anchor.addprevious(paragraph._p)
    ensure_page_setup(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

    if sha256(source) != source_hash_before:
        raise RuntimeError("Source hash changed during build")
    output_blocks = c14n_blocks(output)
    if output_blocks[-len(source_blocks):] != source_blocks:
        raise RuntimeError("Original body blocks are not an identical output suffix")
    output_media = package_manifest(output, "word/media/")
    if not set(source_media.values()).issubset(set(output_media.values())):
        raise RuntimeError("Source media hashes are not a subset of output media hashes")
    output_rels = relationship_manifest(output)
    for rel_id, value in source_rels.items():
        if output_rels.get(rel_id) != value:
            raise RuntimeError(f"Source relationship changed: {rel_id}")

    hashes = {
        "source": source_hash_before,
        "notes": sha256(notes),
        "concepts": sha256(concepts),
        "output": sha256(output),
    }
    print(f"source_sha256={hashes['source']}")
    print(f"notes_sha256={hashes['notes']}")
    print(f"concepts_sha256={hashes['concepts']}")
    print(f"output_sha256={hashes['output']}")
    print(f"source_body_blocks={len(source_blocks)}")
    print(f"output_body_blocks={len(output_blocks)}")
    print("source_body_suffix=IDENTICAL")
    print(f"source_media={len(source_media)}")
    print(f"output_media={len(output_media)}")
    print("source_media_hashes=PRESERVED")
    print("source_relationships=PRESERVED")
    print(f"inserted_paragraphs={len(inserted)}")
    return hashes


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Feishu-import DOCX for course-feishu-pack")
    parser.add_argument("--title", required=True, help="文档标题（主题）")
    parser.add_argument("--notes", type=Path, required=True, help="02 三级笔记 markdown")
    parser.add_argument("--concepts", type=Path, required=True, help="03 概念提取 markdown")
    parser.add_argument("--source", type=Path, required=True, help="00 源文件（.txt 转写或 .docx 原文）")
    parser.add_argument("--output", type=Path, required=True, help="输出 docx 路径")
    parser.add_argument("--subtitle", default="", help="可选副标题行")
    parser.add_argument("--assets-dir", type=Path, default=None, help="图片查找备用目录")
    parser.add_argument("--manifest", type=Path, default=None, help="manifest.json 输出路径（默认与 output 同目录）")
    args = parser.parse_args()

    output = args.output.resolve()
    source = args.source.resolve()
    if output == source:
        raise ValueError("Output must not overwrite the source")
    for name, path in (("notes", args.notes), ("concepts", args.concepts), ("source", args.source)):
        if not path.is_file():
            raise FileNotFoundError(f"{name} not found: {path}")

    if args.source.suffix.lower() == ".docx":
        hashes = build_docx(source, args.notes.resolve(), args.concepts.resolve(), output,
                            args.title, args.subtitle, args.assets_dir)
        source_kind = "docx"
    else:
        hashes = build_txt(source, args.notes.resolve(), args.concepts.resolve(), output,
                           args.title, args.subtitle, args.assets_dir)
        source_kind = "txt"

    manifest_path = args.manifest or output.parent / "manifest.json"
    manifest = {
        "skill": "course-feishu-pack",
        "built_at": __import__("datetime").datetime.now().isoformat(timespec="seconds"),
        "title": args.title,
        "source_kind": source_kind,
        "hashes": hashes,
        "inputs": {
            "notes": str(args.notes),
            "concepts": str(args.concepts),
            "source": str(source),
        },
        "output": str(output),
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"manifest={manifest_path}")


if __name__ == "__main__":
    main()